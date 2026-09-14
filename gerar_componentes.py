"""Gera componentes.json a partir do de-para de combos (planilha da operacao).

Cada produto vendido se decompoe em ate 3 componentes com quantidade por
unidade vendida (ex.: um combo = 2 espetos de carne + 0,2 kg de batata). Isso
permite contar INSUMO (espetos por sabor, kg de batata, litros de chopp) em vez
de contar combos.

PONTE DE NOMES: o PDV renomeou produtos no meio do evento (prefixo "CB"), entao
a planilha usa os nomes antigos e as vendas trazem os novos. Sem a ponte abaixo
ficam 22 produtos / ~16 mil itens (17,9% das vendas) fora da conta - e logo os
campeoes de venda. Cada apelido aponta para o codigo DE equivalente.
"""

import json
import unicodedata

import openpyxl

PLANILHA = r"C:\Users\rfeitosa\Downloads\de-para_RIR (1).xlsx"
SAIDA = "componentes.json"

# nome como sai hoje no PDV -> codigo DE da planilha
APELIDOS = {
    "CB 2 ESPS CARNE + BT":        "2.ESP.CARNE +BT",
    "CB 2 ESPS FRANGO + BT":       "2.ESP.FRANG+BT",
    "CB 1 ESP CARNE + BT":         "ESP.CARNE+BT",
    "CB 1 ESP FRANGO + BT":        "ESP.FRANGO+BT",
    "CB 1 ESP CARNE+1FRANGO+BT":   "1CARN+1FRAN+BT",
    "CB 1 SAND CARNE + BT":        "SAND.CARNE+BT",
    "CB 1 SAND FRANGO + BT":       "SAND.FRANGO+BT",
    "CB MANE BURGUER + BATATA":    "MANE BURGUER+BT",
    "CBAMERICAN FISH + BATATA":    "AMERIC.FISH+BT",
    "COMBO SALAD FISH + BATATA":   "SALAD.FISH+BT",
    "BATATA CRINKLE SIRENE":       "BATATA CRINKLE",
    "CB 1 ESP CARNBTRB COPO":      "1CARN+BT+RB+CP",
    "CB 1 ESP CARN BTRB SEM COPO": "1CARN+BT+RB",
    "CB 1 ESP FRAN+BT+RB COPO":    "1FRANG+BT+RB+CP",
    "CB 1 ESP FRANBTRB SEM COPO":  "1FRAN+BT+RB",
    "CB FISH&CHIPS + RED BULL":    "FISH&CHIP+RB",
    "CB FISH&CHIPS + RED BULL COPO": "FISH&CHIP+RB+CP",
    "CB PIPOCA DE PORCO + 2 CHOP": "CB PIPOCA+2CHOP",
    "CB PIPOCA PORCO + 2 CHOP CP": "PIPOCA+2CHOP CP",
    "CB BT PORQUINHO + RED BULL":  "BAT.PORCO+RB",
    "CB BT PORQUINHO + RB COPO":   "BAT.PORCO+RB+CP",
    "SAND CHURRAS":                "SAND CARNE",
}

# Unidades ausentes na aba "unidades de medidas", deduzidas do uso na planilha.
# BATATA FRITA entra sempre como 0,2 - igual a BATATA CRINKLE, que a propria
# planilha declara em KG. Os demais entram como 1 unidade inteira.
UNIDADES_FALTANTES = {
    "BATATA FRITA": "KG",
    "REDBULL": "UND",
    "AMERICAN FISH": "UND",
    "MANE BURGUER": "UND",
    "SALAD FISH": "UND",
}

# REDBULL generico e RED BULL TRADICIONAL sao o mesmo insumo com dois nomes na
# planilha; unifica para nao dividir a contagem em dois.
SINONIMOS = {"REDBULL": "RED BULL TRADICIONAL"}


def norm(texto):
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(t.upper().split())


def gerar():
    wb = openpyxl.load_workbook(PLANILHA, data_only=True)

    unidades = {}
    ws = wb["unidades de medidas"]
    for i in range(2, ws.max_row + 1):
        nome, un = ws.cell(i, 1).value, ws.cell(i, 2).value
        if nome:
            unidades[norm(nome)] = (str(un).strip().upper() if un else "UND")
    for nome, un in UNIDADES_FALTANTES.items():
        unidades.setdefault(norm(nome), un)

    produtos, rotulos = {}, {}
    ws = wb["Vendas"]
    for i in range(2, ws.max_row + 1):
        de = ws.cell(i, 1).value
        if not de:
            continue
        comps = []
        for col in (4, 6, 8):
            nome, qtd = ws.cell(i, col).value, ws.cell(i, col + 1).value
            if nome and qtd:
                alvo = SINONIMOS.get(norm(nome), norm(nome))
                # Guarda a grafia ORIGINAL (com acento) para exibir: a chave e
                # normalizada sem acento so para casar com o nome do PDV. So
                # quando NAO houve sinonimo - senao "REDBULL" viraria o rotulo
                # de "RED BULL TRADICIONAL".
                if norm(nome) == alvo:
                    rotulos.setdefault(alvo, str(nome).strip())
                comps.append([alvo, float(qtd)])
        # Produto sem decomposicao na planilha (ex.: AGUA CRYSTAL) conta como
        # 1 unidade de si mesmo - senao ele simplesmente sumiria da contagem.
        if not comps:
            comps = [[norm(de), 1.0]]
            unidades.setdefault(norm(de), "UND")
        produtos[norm(de)] = {"marca": str(ws.cell(i, 3).value or "").strip(),
                              "componentes": comps}

    # aplica a ponte de nomes
    for apelido, destino in APELIDOS.items():
        alvo = produtos.get(norm(destino))
        if alvo is None:
            raise SystemExit(f"apelido sem destino no de-para: {apelido} -> {destino}")
        produtos[norm(apelido)] = alvo

    # sinonimos tambem valem para as unidades
    for origem, destino in SINONIMOS.items():
        unidades.pop(origem, None)
        unidades.setdefault(destino, "UND")

    # Categoria de cada COMPONENTE (a planilha so traz a marca do PRODUTO).
    # Um componente e bebida quando aparece em algum produto cuja marca comeca
    # por "Bebida" - assim o Red Bull que vem dentro do combo do Espetto nao
    # entra no insight de COMIDA da marca.
    categorias = {}
    for info in produtos.values():
        bebida = info["marca"].lower().startswith("bebida")
        for nome, _ in info["componentes"]:
            if bebida or nome not in categorias:
                categorias[nome] = "Bebida" if bebida else "Comida"

    # Subcategoria da bebida (CHOPE, REDBULL, SOFT DRINK...), tirada da marca
    # do produto que a vende pura. Assim o chopp que vem dentro do combo do
    # Mane tambem cai no grupo CHOPE, e nao no do Mane.
    subcategorias = {}
    for info in produtos.values():
        marca = info["marca"]
        if not marca.lower().startswith("bebida"):
            continue
        grupo = marca.split("/")[-1].strip().upper()
        for nome, _ in info["componentes"]:
            subcategorias.setdefault(nome, grupo)

    with open(SAIDA, "w", encoding="utf-8") as saida:
        json.dump({"unidades": unidades, "categorias": categorias,
                   "subcategorias": subcategorias, "rotulos": rotulos,
                   "produtos": produtos},
                  saida, ensure_ascii=False, indent=1)
    return produtos, unidades, categorias


if __name__ == "__main__":
    p, u, c = gerar()
    print(f"{SAIDA}: {len(p)} produtos, {len(u)} componentes")
    print("  comida:", sorted(k for k, v in c.items() if v == "Comida"))
    print("  bebida:", sorted(k for k, v in c.items() if v == "Bebida"))
